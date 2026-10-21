from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import xlrd
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from werkzeug.datastructures import FileStorage
from werkzeug.utils import secure_filename

from .sql_analyzer import analyze_sql, report_as_prompt


MAX_EXTRACTED_CHARS = 120_000
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
ALLOWED_EXTENSIONS = {".pdf", ".sql", ".xlsx", ".xls", ".docx", ".txt", ".csv"} | IMAGE_EXTENSIONS


class FileExtractionError(ValueError):
    pass


@dataclass
class FileExtraction:
    filename: str
    text: str
    mime_type: str | None = None
    binary_data: bytes | None = None
    sql_report: dict | None = None


def _truncate(text):
    text = text.strip()
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    return text[:MAX_EXTRACTED_CHARS] + "\n\n[Contenu tronqué pour respecter la limite de contexte.]"


def _decode_text(data):
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileExtractionError("Impossible de décoder ce fichier texte.")


def _extract_pdf(data):
    reader = PdfReader(BytesIO(data))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _extract_docx(data):
    document = Document(BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(paragraphs)


def _extract_xlsx(data):
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"## Feuille: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append("\t".join(values))
    return "\n".join(lines)


def _extract_xls(data):
    workbook = xlrd.open_workbook(file_contents=data)
    lines = []
    for sheet in workbook.sheets():
        lines.append(f"## Feuille: {sheet.name}")
        for row_index in range(sheet.nrows):
            values = [str(value) for value in sheet.row_values(row_index)]
            if any(value.strip() for value in values):
                lines.append("\t".join(values))
    return "\n".join(lines)


def extract_file(file: FileStorage):
    if not file or not file.filename:
        return None

    filename = secure_filename(file.filename)
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise FileExtractionError(
            "Format non pris en charge. Utilisez PDF, image, SQL, Excel, Word, TXT ou CSV."
        )

    data = file.read()
    if not data:
        raise FileExtractionError("Le fichier importé est vide.")

    if extension in IMAGE_EXTENSIONS:
        return FileExtraction(
            filename=filename,
            text="Image jointe à analyser visuellement.",
            mime_type=file.mimetype or f"image/{extension.lstrip('.')}",
            binary_data=data,
        )

    if extension == ".pdf":
        text = _extract_pdf(data)
        if not text:
            return FileExtraction(
                filename=filename,
                text="PDF scanné ou sans texte extractible, joint à Gemini pour analyse.",
                mime_type="application/pdf",
                binary_data=data,
            )
    elif extension == ".docx":
        text = _extract_docx(data)
    elif extension == ".xlsx":
        text = _extract_xlsx(data)
    elif extension == ".xls":
        text = _extract_xls(data)
    else:
        text = _decode_text(data)

    sql_report = analyze_sql(text) if extension == ".sql" else None
    if sql_report:
        text = f"{text}\n\n---\n{report_as_prompt(sql_report)}"

    return FileExtraction(filename=filename, text=_truncate(text), sql_report=sql_report)
